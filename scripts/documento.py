"""Genera el documento de postulación (DOCX) con las cifras de web/public/datos, para que web y PDF no se contradigan."""
import io
import json
import math
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

AUTORIA = sys.argv[1] if len(sys.argv) > 1 else "[Nombre de la autora]"
WEB = sys.argv[2] if len(sys.argv) > 2 else "[enlace de la web]"
REPO = sys.argv[3] if len(sys.argv) > 3 else "[enlace del repositorio]"
REVISION = sys.argv[4] if len(sys.argv) > 4 else ""

D = "web/public/datos"
r = json.load(open(f"{D}/resumen.json"))
com = {c["slug"]: json.load(open(f"{D}/comunas/{c['slug']}.json")) for c in r["comunas"]}
cr = r["crecimiento"]
po = r["poblacion"]
hist = r["sectores_historial"]


def partes(o):
    t = o["ya_estaba_ha"] + o["llego_ciudad_ha"] + o["plantacion_nueva_ha"]
    return {k: o[k] / t * 100 for k in ("ya_estaba_ha", "llego_ciudad_ha", "plantacion_nueva_ha")}


casos = [partes(cr[c][d]) for c in ("col1_2022", "col2_2022") for d in ("100m", "500m")]
rango = lambda k: (min(x[k] for x in casos), max(x[k] for x in casos))
c1, c2 = r["nacional"]["col1"], r["nacional"]["col2"]
v1, v2 = r["validacion"]["col1"], r["validacion"]["col2"]
sens = r["sensibilidad"]["col1"]
sin_una = r["sin_una_comuna"]["col1"]
n_sect = r["sectores_total"]
suma_sect = r["sectores_suma_100m_ha"]
top_sect = r["sectores_top"]


def n0(x):
    # Redondeo hacia arriba en .5, igual que Math.round de la web
    return f"{math.floor(x + 0.5):,}".replace(",", ".")


def n2(x):
    return f"{x:.2f}".replace(".", ",")


def n1(x):
    return f"{x:.1f}".replace(".", ",")


def nombre_sector(x):
    if x["lugar_m"] <= 500:
        return x["lugar"]
    if x["lugar_m"] <= 2000:
        d = f"{math.floor(x['lugar_m'] / 10 + 0.5) * 10} m" if x["lugar_m"] < 1000 else f"{x['lugar_m'] / 1000:.1f} km".replace(".", ",")
        return f"Cerca de {x['lugar']} (a {d})"
    return "Sin lugar con nombre a menos de 2 km"


def pct(a, b):
    return f"{(b - a) / a * 100:+.0f} %".replace("+", "+").replace("-", "−")


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Aptos"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.2
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.3)
    s.top_margin = s.bottom_margin = Cm(2.2)
for nivel, tam in ((1, 15), (2, 12.5)):
    h = doc.styles[f"Heading {nivel}"]
    h.font.name = "Aptos"
    h.font.size = Pt(tam)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x16, 0x18, 0x1A)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True  # un título nunca queda solo al final de una página


def p(texto, *, italic=False, size=None, align=None, color=None):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if align:
        par.alignment = align
    return par


def lista(items):
    pars = [doc.add_paragraph(t, style="List Bullet") for t in items]
    # El último punto nunca queda solo al comienzo de una página
    if len(pars) > 1:
        pars[-2].paragraph_format.keep_with_next = True


def _fila_entera(row, encabezado=False):
    """La fila no se corta entre páginas; la de encabezado se repite si la tabla continúa."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if encabezado:
        trPr.append(OxmlElement("w:tblHeader"))


def tabla(cabecera, filas, anchos=None, pie=None):
    # La frase que presenta la tabla queda en la misma página que ella
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = "Table Grid"
    for i, h in enumerate(cabecera):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for f in filas:
        cells = t.add_row().cells
        for i, v in enumerate(f):
            cells[i].text = str(v)
            # Solo las cifras van a la derecha
            if i > 0 and any(ch.isdigit() for ch in str(v)[:3]):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                for run in par.runs:
                    run.font.size = Pt(9)
    if anchos:
        for row in t.rows:
            for i, w in enumerate(anchos):
                row.cells[i].width = Cm(w)
    # Tabla junta: cada fila se mantiene con la siguiente (y la última con su pie, si lo tiene)
    for k, row in enumerate(t.rows):
        _fila_entera(row, encabezado=k == 0)
        if k < len(t.rows) - 1 or pie:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    if pie:
        pp = p(pie, size=8.5, color="55595E")
        pp.paragraph_format.space_before = Pt(3)
        pp.paragraph_format.space_after = Pt(10)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def imagen(slug, ancho_cm, pie):
    buf = io.BytesIO()
    Image.open(f"web/public/mapas/{slug}.webp").convert("RGB").save(buf, "PNG")
    buf.seek(0)
    doc.add_picture(buf, width=Cm(ancho_cm))
    doc.paragraphs[-1].paragraph_format.keep_with_next = True  # el mapa con su pie
    p(pie, size=8.5, color="55595E")


# Portada
p("Premio MapBiomas Chile, primera edición (2026)", size=9, color="55595E")
p("Categoría 2: Aplicaciones generales, políticas públicas e iniciativas educativas o sociales", size=9, color="55595E")
t = doc.add_paragraph()
run = t.add_run("Límite Vivo")
run.bold = True
run.font.size = Pt(24)
p("Plantaciones forestales en el borde de las ciudades de Chile: cuánta hay, quién se acercó a quién y quiénes viven al lado, por comuna y por sector, 2000 a 2024", size=13)
p(f"Autoría: {AUTORIA}", size=10)
p(f"Herramienta en línea: {WEB}", size=10)
p(f"Código y datos: {REPO}", size=10)

doc.add_heading("Resumen", 1)
p(
    f"Límite Vivo es una herramienta pública que calcula, para cada una de las 343 comunas de Chile, cuántas hectáreas de "
    f"plantación forestal hay a menos de 100, 200 y 500 metros de una zona poblada, cómo cambió esa franja desde el año 2000, "
    f"si cambió porque la ciudad creció hacia la plantación o porque se plantó junto a las casas, cuántas personas viven al lado "
    f"y qué pasó con esa franja en los incendios. El cálculo se basa en los mapas anuales de cobertura y uso del suelo de "
    f"MapBiomas Chile (Colecciones 1 y 2) y en MapBiomas Fuego (Colección 1); el Censo 2024 del INE se usa como complemento para "
    f"contar a las personas."
)
p(
    f"En 2022 había {n0(c1['2022']['plantacion_500m_ha'])} ha de plantación a menos de 500 m de una zona poblada y "
    f"{n0(c1['2022']['plantacion_100m_ha'])} ha a menos de 100 m (Colección 1). La franja de 500 m creció "
    f"{pct(c1['2000']['plantacion_500m_ha'], c1['2022']['plantacion_500m_ha'])} entre 2000 y 2022 según la Colección 1 y "
    f"{pct(c2['2000']['plantacion_500m_ha'], c2['2024']['plantacion_500m_ha'])} entre 2000 y 2024 según la Colección 2. La "
    f"diferencia se explica por cómo cada colección clasifica la infraestructura, y por eso se informan ambas."
)
p(
    f"En las comunas con incendios cerca de zonas pobladas, la plantación de la franja de 500 m se quemó en mayor proporción que "
    f"la vegetación natural de la misma franja, {n2(v2['natural']['rr'])} veces más con la Colección 2 (intervalo de 95 %, "
    f"{n2(v2['natural']['ic95'][0])} a {n2(v2['natural']['ic95'][1])}) y {n2(v1['natural']['rr'])} veces más con la Colección 1 "
    f"({n2(v1['natural']['ic95'][0])} a {n2(v1['natural']['ic95'][1])}). Es una asociación observada en tres temporadas, no una "
    f"relación causal, y tiene excepciones como Viña del Mar en 2024."
)
lc1, lc2 = rango("llego_ciudad_ha"), rango("plantacion_nueva_ha")
p(
    f"La franja no creció por un solo motivo. De la plantación que en 2022 estaba a menos de 100 m de las casas, "
    f"{n0(cr['col1_2022']['100m']['llego_ciudad_ha'])} ha ya existían en 2000 más lejos y fue la zona poblada la que creció hasta "
    f"ellas, y {n0(cr['col1_2022']['100m']['plantacion_nueva_ha'])} ha se plantaron después de 2005 (Colección 1). Según la colección y "
    f"la distancia, entre el {n1(lc1[0])} % y el {n1(lc1[1])} % de la plantación junto a las casas corresponde al crecimiento de la "
    f"ciudad y entre el {n1(lc2[0])} % y el {n1(lc2[1])} % a plantación nueva. Con las manzanas del Censo 2024, entre "
    f"{n0(po['centro']['100m'])} y {n0(po['100m']['personas'])} personas viven a menos de 100 m de plantación forestal, según si se "
    f"mide desde el centro o desde el borde de su manzana."
)
p(
    f"Como una comuna puede tener la plantación concentrada en pocos bordes, el borde de cada zona poblada se divide además en "
    f"sectores de alrededor de 1 km. Hay {n0(n_sect)} sectores con plantación a menos de 100 m, cada uno con su ubicación y el nombre "
    f"del barrio, población o localidad más cercana según OpenStreetMap. "
    f"Para cada sector se informa además cuánto de su franja de 100 m se quemó entre 2013 y 2025. "
    "La herramienta entrega una ficha por comuna con mapa, sectores, cifras por distancia, evolución, origen de la franja, personas que viven al lado, contraste con incendios y usos "
    "posibles en el Plan Comunal para la Reducción del Riesgo de Desastres (Ley 21.364). Los datos se descargan en CSV y el "
    "código es público."
)

doc.add_heading("1. Problema", 1)
p(
    "Los incendios de 2023 en Ñuble, Biobío y La Araucanía y el de febrero de 2024 en la Región de Valparaíso mostraron el daño "
    "que produce el fuego cuando llega a zonas pobladas. Una parte de ese riesgo está en el límite entre las ciudades y las "
    "plantaciones de pino y eucalipto, que en muchas comunas llegan hasta el borde de las casas."
)
p(
    "Ese límite se puede medir con los mapas de MapBiomas, pero no encontramos una cifra comparable por comuna de cuánta "
    "plantación hay junto a las zonas pobladas que un municipio pueda usar directamente. La Ley 21.364 (artículo 28) exige a cada comuna un Plan Comunal para la Reducción del Riesgo de Desastres, y "
    "según Corporación Ciudades, con datos de SENAPRED al 6 de julio de 2026, 127 de 345 comunas no lo tenían. El proyecto de "
    "ley de incendios forestales y rurales (boletín 16.335-14), que propone zonas de interfaz y de amortiguación, seguía en "
    "trámite en octubre de 2026."
)

doc.add_heading("2. Antecedentes", 1)
p(
    "La relación entre plantaciones e incendios en Chile está documentada. McWethy et al. (2018) mostraron que entre 2001 y "
    "2017 el fuego ocurrió de preferencia en plantaciones exóticas. Jaque Castillo et al. (2021) identificaron la cercanía entre "
    "la expansión urbana y las plantaciones como el factor principal del riesgo en el Gran Concepción. Hora et al. (2025) "
    "encontraron más plantación a menos de 30 m de las viviendas destruidas en el incendio de 2023 en el Biobío. Al presentar "
    "MapBiomas Fuego, sus autores informaron que cerca del 30 % de los incendios afectó áreas próximas a ciudades."
)
p(
    "Existen también mapas nacionales de interfaz urbano-forestal. El más cercano es el análisis nacional de riesgo y zonas de "
    "interfaz de CONAF (2021, actualizado en 2024), de 10 m de resolución y con un visor público, que es la referencia oficial "
    "para conocer el riesgo de un lugar. Según su descripción, no separa la plantación forestal de la vegetación nativa ni "
    "compara años. Hay además mapas de interfaz para todo el país en Schug et al. (2023) y Gelabert et al. (2026)."
)
p(
    "Sarricolea et al. (2020) estimaron que cerca de 3 millones de personas viven en la interfaz urbano-forestal entre "
    "Valparaíso y La Araucanía, con resultados por región y considerando toda la vegetación inflamable. Hernández et al. (2024) "
    "midieron con MapBiomas el crecimiento de la infraestructura urbana de Chile entre 2000 y 2022, por región. El Ministerio de "
    "Educación (visor de comunidades educativas ante amenazas socionaturales) y SENAPRED (Visor Chile Preparado) ya muestran "
    "colegios, jardines infantiles y centros de salud sobre la amenaza de incendio de CONAF, por lo que Límite Vivo no repite esa "
    "capa."
)
p(
    "En nuestra revisión no encontramos una herramienta pública que, para cada comuna y con las series de MapBiomas, cuantifique "
    "la plantación forestal a distancias fijas de las zonas pobladas, la siga desde el año 2000, separe si fue la ciudad la que "
    "llegó a la plantación o la plantación la que llegó a la ciudad, cuente a las personas que viven al lado y contraste todo con "
    "el área quemada dentro de la misma comuna y franja. Límite Vivo se propone como complemento de esos trabajos, en particular "
    "del análisis de CONAF, y lo enlaza desde cada ficha comunal y desde la búsqueda por dirección."
)

doc.add_heading("3. Objetivo", 1)
p(
    "Entregar a municipios, organismos de emergencia y a la ciudadanía un indicador comunal, reproducible y actualizable con "
    "cada colección de MapBiomas, de cuánta plantación forestal hay junto a las zonas pobladas, cómo ha cambiado, si ese cambio "
    "viene del crecimiento urbano o de plantación nueva, cuántas personas viven al lado y cómo se comportó esa franja en los "
    "incendios recientes."
)

doc.add_heading("4. Datos", 1)
lista([
    "MapBiomas Chile, cobertura y uso del suelo, Colección 1 (2000 y 2022, y 2000 a 2005 para identificar la plantación existente). Base de las cifras principales. Se leyó del asset público en Google Earth Engine, en su grilla nativa de 30 m, y da exactamente las mismas cifras que los GeoTIFF anuales del almacenamiento público, que se usaron primero y que el 30 de octubre de 2026 dejaron de estar disponibles para descarga directa.",
    "MapBiomas Chile, cobertura y uso del suelo, Colección 2 (2000 a 2005, 2022 y 2024), leída en Google Earth Engine y descargada en la misma grilla. Análisis de sensibilidad.",
    "MapBiomas Fuego Chile, Colección 1: área quemada anual 2023, 2024 y 2025, y frecuencia de quema 2013-2025. Cubre de Valparaíso a La Araucanía.",
    "Límites comunales de MapBiomas Chile (nivel político 4, 343 comunas), publicados como ráster en Earth Engine, vectorizados y nombrados por superposición con la capa oficial de comunas del MOP (servicio DPA); cada comuna coincide en más del 92 % de su superficie. Estadísticas oficiales por comuna de ambas colecciones.",
    "Complemento: cartografía del Censo de Población y Vivienda 2024 (INE), manzanas urbanas y de aldeas con personas, viviendas ocupadas, grupos de edad y personas con mucha dificultad para caminar o subir escaleras.",
])

doc.add_heading("5. Método", 1)
lista([
    "Zonas pobladas. Se usa la clase infraestructura urbana (código 24). Como incluye carreteras principales (Hernández et al., 2024), se eliminan las formas de menos de 3 píxeles de ancho con una apertura morfológica y las manchas de menos de 2 ha.",
    "Distancia. Para cada píxel se calcula la distancia en metros hasta el borde de la zona poblada más cercana, con el tamaño real del píxel según la latitud. Se consideran zonas pobladas hasta cerca de 1 km fuera de la comuna.",
    "Franjas. Se suman las hectáreas de plantación forestal (código 9) dentro de la comuna a menos de 100, 200 y 500 m, con el área de cada píxel según su latitud.",
    "Sectores. Con la cobertura de 2022 se marca el borde de cada zona poblada que tiene plantación a menos de 100 m. Los tramos expuestos a menos de unos 90 m entre sí se unen y se cortan en celdas de 0,01° (cerca de 1 km). Cada píxel de plantación a menos de 100 m se asigna al punto de zona poblada más cercano para sumar las hectáreas de cada sector. El largo del borde es aproximado (se cuenta en píxeles). Cada sector toma el nombre del lugar más cercano de OpenStreetMap (barrio, población, villa o localidad); si está a más de 500 m se indica la distancia y si está a más de 2 km se omite el nombre.",
    "Quién se acercó a quién. La plantación que en 2022 está a menos de 100 o 500 m de una zona poblada se separa, píxel a píxel, en tres partes que suman exacto: la que ya estaba junto a las casas en 2000; la que existía pero más lejos y quedó cerca porque la zona poblada creció hasta ella; y la que se plantó después. Una plantación se considera existente si aparece como plantación en algún año entre 2000 y 2005, para no contar la cosecha y la replantación como plantación nueva. Se calcula con ambas colecciones.",
    "Personas. Cada manzana del Censo 2024 se lleva a la grilla de MapBiomas y se mide su distancia mínima al borde de la plantación de 2022. Si alguna parte de la manzana está dentro de la distancia, cuentan todas sus personas (criterio amplio); como control se cuenta también la manzana solo si su punto interior está dentro (criterio estricto), y se informa el rango.",
    "Historial de quema por sector. Para la franja de 100 m junto a cada sector, cualquiera sea su cobertura, se suma la superficie quemada al menos una vez entre 2013 y 2025 y la quemada dos o más veces, con la frecuencia de quema de MapBiomas Fuego.",
    "Contraste con incendios. Con la cobertura de 2022 y el área quemada 2023 a 2025, se compara dentro de cada comuna (para no mezclar lugares con distinto clima, relieve e incendios) la proporción quemada de la plantación de la franja de 500 m con la de la vegetación natural y la agricultura de la misma franja. Las comunas se combinan con la razón de Mantel-Haenszel y el intervalo de confianza de 95 % se obtiene con 2.000 remuestreos de comunas. Entran las comunas con al menos 20 ha de plantación en la franja y 10 ha quemadas en ella.",
])

doc.add_heading("6. Control de calidad", 1)
lista([
    "En 960 comparaciones de plantación e infraestructura por comuna contra las estadísticas oficiales de la Colección 1, la diferencia mediana fue de 0,04 % y la máxima de 0,31 %.",
    "Con la Colección 2 la plantación concilia (mediana de 0,07 %), pero la infraestructura difiere cerca de 3 % de forma sistemática. La imagen de Earth Engine no es exactamente la versión de las estadísticas publicadas, y por eso esa colección se usa como análisis de sensibilidad.",
    "Las franjas de seis comunas se recalcularon con un método independiente (polígonos, buffer en UTM y área elipsoidal). Las diferencias fueron menores a 1 % en la franja de 500 m y a 3 % en la de 200 m, y llegaron a 9 % en la de 100 m; en las seis el método principal dio la cifra más baja, así que la franja de 100 m tiende a quedarse corta.",
    f"La suma de la plantación de los {n0(n_sect)} sectores ({n0(suma_sect)} ha) difiere en {n1(abs(suma_sect / c1['2022']['plantacion_100m_ha'] - 1) * 100)} % de la cifra nacional a menos de 100 m. Ninguna comuna con 5 ha o más a esa distancia queda sin sectores, y su ubicación se revisó en mapas de control.",
    "En la búsqueda por dirección, la distancia de 12 celdas al azar se recalculó directamente sobre el ráster, con una diferencia máxima de 2 m.",
    "Hay pruebas automáticas del filtro de zonas pobladas, del cálculo de Mantel-Haenszel, de la descomposición por origen (que las partes sumen exacto y que un rodal replantado no cuente como nuevo) y del conteo de personas, y las comunas sin plantación dan cero.",
    "En cada comuna, las tres partes del origen suman exactamente la franja del cálculo principal. Todas las manzanas del censo quedaron medidas; solo faltan Juan Fernández e Isla de Pascua, que no están entre las 343 comunas de MapBiomas. En los sectores, la superficie quemada entre 2013 y 2025 nunca es menor que la plantación quemada entre 2023 y 2025 de la misma franja.",
    "El control encontró dos errores, que se corrigieron antes de entregar. Los mapas de control mostraron que la infraestructura incluía carreteras, y la verificación independiente mostró que la primera versión medía la distancia de centro a centro de píxel, lo que subestimaba la franja de 100 m en cerca de 15 %.",
])

doc.add_heading("7. Resultados nacionales", 1)
tabla(
    ["Distancia a zonas pobladas", "2000 (C1)", "2022 (C1)", "2000 (C2)", "2024 (C2)"],
    [[t, n0(c1["2000"][k]), n0(c1["2022"][k]), n0(c2["2000"][k]), n0(c2["2024"][k])] for k, t in (
        ("plantacion_100m_ha", "A menos de 100 m"), ("plantacion_200m_ha", "A menos de 200 m"), ("plantacion_500m_ha", "A menos de 500 m"))],
    pie="Hectáreas de plantación forestal. C1 y C2: Colecciones 1 y 2 de MapBiomas Chile.",
)
p(
    f"Leídas como franjas de protección, una franja de 100 m alrededor de las zonas pobladas abarcaría hoy "
    f"{n0(c1['2022']['plantacion_100m_ha'])} ha de plantación y una de 500 m, {n0(c1['2022']['plantacion_500m_ha'])} ha (Colección 1, "
    f"2022). La herramienta entrega esa cifra para cada comuna, lo que permite estimar cuánta plantación quedaría dentro de una norma de distancia en cada lugar."
)
top = [c for c in r["comunas"] if c["puesto"]][:15]
p(f"Las 15 comunas con más plantación a menos de 500 m de zonas pobladas en 2022 (Colección 1):")
tabla(
    ["Comuna", "Región", "500 m, 2000", "500 m, 2022", "100 m, 2022"],
    [[f"{c['puesto']}. {c['comuna']}", c["region"], n0(c["p500_2000"]), n0(c["p500_2022"]), n0(c["p100_2022"])] for c in top],
)

p("Los 10 sectores del país con más plantación forestal a menos de 100 m de una zona poblada (cobertura 2022, Colección 1):")
tabla(
    ["Sector", "Comuna", "Plantación a menos de 100 m (ha)", "Borde expuesto aprox. (m)"],
    [[nombre_sector(x), f"{x['comuna']} (sector {x['n']})", n0(x["plantacion_100m_ha"]), n0(math.floor(x["borde_m"] / 10 + 0.5) * 10)] for x in top_sect[:10]],
)

doc.add_heading("Quién se acercó a quién", 2)
p(
    "La plantación que hoy está junto a las casas tiene tres orígenes, que se resuelven en lugares distintos. La que ya existía y "
    "a la que llegó la ciudad depende del Plan Regulador Comunal y de los permisos de loteo; la plantación nueva junto a las casas "
    "depende de dónde se planta y de las franjas que se exijan a los predios forestales."
)
tabla(
    ["Caso", "Ya estaba (ha)", "Llegó la ciudad (ha)", "Plantación nueva (ha)"],
    [[t, n0(cr[c][d]["ya_estaba_ha"]), n0(cr[c][d]["llego_ciudad_ha"]), n0(cr[c][d]["plantacion_nueva_ha"])] for t, c, d in (
        ("C1, 2000-2022, a menos de 100 m", "col1_2022", "100m"), ("C1, 2000-2022, a menos de 500 m", "col1_2022", "500m"),
        ("C2, 2000-2022, a menos de 100 m", "col2_2022", "100m"), ("C2, 2000-2022, a menos de 500 m", "col2_2022", "500m"),
        ("C2, 2000-2024, a menos de 100 m", "col2_2024", "100m"))],
)
p(
    f"Las colecciones no dan las mismas proporciones: el crecimiento de la ciudad explica entre el {n1(lc1[0])} % y el {n1(lc1[1])} % "
    f"y la plantación nueva entre el {n1(lc2[0])} % y el {n1(lc2[1])} %. Ambos caminos pesan en las dos colecciones, y la ficha de cada "
    f"comuna indica cuál domina en ella. Por ejemplo, en Tomé "
    f"{n0(com['tome']['crecimiento']['100m']['llego_ciudad_ha'])} de sus {n0(com['tome']['crecimiento']['100m']['final_ha'])} ha a menos de "
    f"100 m corresponden a plantación que ya existía y a la que llegó la ciudad."
)

doc.add_heading("Personas que viven al lado", 2)
p(
    f"Según las manzanas del Censo 2024, entre {n0(po['centro']['100m'])} y {n0(po['100m']['personas'])} personas viven a menos de "
    f"100 m de plantación forestal (criterio estricto y amplio) y entre {n0(po['centro']['500m'])} y {n0(po['500m']['personas'])} a "
    f"menos de 500 m. En las manzanas que llegan a menos de 100 m hay {n0(po['100m']['viviendas'])} viviendas ocupadas, "
    f"{n0(po['100m']['ninos_0_13'])} niñas y niños de 0 a 13 años, {n0(po['100m']['mayores_60'])} personas de 60 años o más y "
    f"{n0(po['100m']['dificultad_mover'])} personas que no pueden caminar o subir escaleras, o lo hacen con mucha dificultad, un dato "
    f"directo para los planes de evacuación. A diferencia de Sarricolea et al. (2020), la cifra considera solo plantación forestal, "
    f"cubre todo el país y se entrega por comuna."
)
tabla(
    ["Comuna", "Región", "Manzanas a menos de 100 m (personas)", "A menos de 500 m", "Parte de quienes viven en manzanas"],
    [[x["comuna"], x["region"], n0(x["personas"]), n0(x["personas_500m"]), f"{n1(x['personas'] / x['personas_total'] * 100)} %"]
     for x in r["top_personas"][:10]],
    pie="Las 10 comunas con más personas en manzanas que llegan a menos de 100 m de plantación (criterio amplio).",
)

doc.add_heading("Historial de quema de los sectores", 2)
p(
    f"De {n0(hist['con_dato'])} sectores con datos de MapBiomas Fuego, {n0(hist['quemados'])} tuvieron al menos media hectárea de su "
    f"franja de 100 m quemada entre 2013 y 2025 y {n0(hist['reincidentes'])} tuvieron media hectárea o más quemada dos o más veces. "
    f"La ficha de cada comuna muestra el dato sector por sector."
)

doc.add_heading("8. Contraste con los incendios de 2023 a 2025", 1)
tabla(
    ["Plantación frente a", "Colección 1", "Colección 2"],
    [[nombre,
      f"{n2(v1[k]['rr'])} ({n2(v1[k]['ic95'][0])} a {n2(v1[k]['ic95'][1])}); {v1[k]['comunas_mayor']} de {v1[k]['comunas']} comunas",
      f"{n2(v2[k]['rr'])} ({n2(v2[k]['ic95'][0])} a {n2(v2[k]['ic95'][1])}); {v2[k]['comunas_mayor']} de {v2[k]['comunas']} comunas"]
     for k, nombre in (("natural", "Vegetación natural"), ("agricola", "Agricultura y pasturas"))],
)
p(
    f"Con ambas colecciones la plantación cercana a zonas pobladas se quemó en mayor proporción que la vegetación natural de la "
    f"misma franja. Con la Colección 2 el intervalo queda sobre 1; con la Colección 1 el límite inferior llega a "
    f"{n2(v1['natural']['ic95'][0])}, de modo que esa colección por sí sola no es concluyente. El resultado se mantiene al cambiar "
    f"los umbrales de inclusión (entre {n2(min(x['natural'] for x in sens))} y {n2(max(x['natural'] for x in sens))} veces) y al "
    f"sacar una comuna a la vez (entre {n2(sin_una[0])} y {n2(sin_una[1])} veces)."
)
vi = com["vina-del-mar"]["fuego_col1"]
to = com["tome"]["fuego_col1"]
p(
    f"Casos. En Tomé se quemó el {n1(to['plantacion_quemada_ha'] / to['plantacion_ha'] * 100)} % de la plantación de la franja y "
    f"el {n1(to['natural_quemada_ha'] / to['natural_ha'] * 100)} % de la vegetación natural. En Viña del Mar ocurrió lo contrario, con el "
    f"{n1(vi['plantacion_quemada_ha'] / vi['plantacion_ha'] * 100)} % de la plantación quemada frente al "
    f"{n1(vi['natural_quemada_ha'] / vi['natural_ha'] * 100)} % de la vegetación natural. Según estos mapas, el incendio de 2024 "
    f"quemó en proporción más matorral y bosque que plantación, y la herramienta muestra ese caso igual que los demás."
)
imagen("tome", 9.5, "Tomé. Rojo: zona poblada. Naranjo: plantación a menos de 500 m. Café: otra plantación. Verde claro: vegetación natural. Zonas oscurecidas: quemado 2023 a 2025. Cobertura 2022, Colección 1.")

doc.add_heading("9. La herramienta", 1)
lista([
    "Inicio con las cifras nacionales, el contraste con incendios, un mapa clicable de las comunas de Valparaíso a Los Lagos, los 25 sectores con más plantación junto a zonas pobladas y un ranking filtrable por región.",
    "Una ficha por comuna (343) con mapa, sectores numerados con enlace a OpenStreetMap y su historial de quema 2013-2025, comunas vecinas, tabla por distancia en ambas colecciones, indicación de si las colecciones coinciden en la tendencia, origen de la franja (quién se acercó a quién), personas que viven al lado según el Censo 2024, contraste con incendios cuando hay cobertura de Fuego y usos para el municipio.",
    "Búsqueda por dirección con sugerencias mientras se escribe (Photon, datos de OpenStreetMap) o con la ubicación del navegador, y un mapa opcional para ajustar el punto cuando OpenStreetMap no tiene el número de la casa. Muestra a cuántos metros está la plantación más cercana, cuántas hectáreas hay en 500 m y si hubo área quemada cerca entre 2023 y 2025. La ubicación del navegador no se envía a ningún servicio externo.",
    "Enlace al visor oficial de riesgo de CONAF desde cada ficha y desde la búsqueda por dirección.",
    "Metodología completa con antecedentes, control de calidad y limitaciones.",
    "Datos descargables en CSV y código abierto en Python (rasterio, numpy, scipy, geopandas) y en la web (Next.js, exportación estática).",
])

doc.add_heading("10. Impacto potencial", 1)
p(
    "Un municipio puede incorporar la ficha como antecedente de amenaza de incendio forestal en su Plan Comunal para la "
    "Reducción del Riesgo de Desastres, priorizar cortafuegos, limpieza de bordes y rutas de evacuación donde la plantación está "
    "a menos de 100 m de las casas, saber cuántas personas mayores, niños y personas con dificultad para moverse viven en esas "
    "manzanas al planificar evacuaciones, y usar la cifra al revisar el Plan Regulador Comunal sabiendo cuánto de la franja viene "
    "del crecimiento de la propia ciudad. Los organismos de emergencia ganan una "
    "vista nacional comparable. Si se aprueba una ley con zonas de amortiguación, la herramienta permite estimar cuánta plantación "
    "queda dentro de franjas de 100, 200 o 500 m en cada comuna. Y las organizaciones vecinales y la prensa tienen un dato verificable de su propia comuna. "
    "Como se apoya solo en datos abiertos de MapBiomas, se puede actualizar con cada nueva colección. La temporada de incendios "
    "2026-2027 coincide además con un El Niño fuerte a muy fuerte según el Centro de Predicción Climática de la NOAA (8 de octubre "
    "de 2026), lo que hace oportuno contar con estos datos antes del verano."
)

doc.add_heading("11. Limitaciones", 1)
lista([
    "Mide exposición y no probabilidad de incendio, porque no considera igniciones, viento, pendiente, sequía ni manejo.",
    "Depende de la clasificación de MapBiomas; las colecciones difieren sobre todo en infraestructura, lo que cambia la magnitud de la tendencia.",
    "Con píxeles de 30 m, la franja de 100 m tiene una incertidumbre de método de hasta 10 %. No incluye casas aisladas ni zonas pobladas de menos de 2 ha.",
    "El contraste con incendios cubre tres temporadas y solo de Valparaíso a La Araucanía, y muestra asociación, no causalidad.",
    "El origen de la franja depende de que la clasificación de zonas pobladas y plantación sea comparable entre años; como las colecciones dan proporciones distintas, se informa un rango.",
    "Las personas son las de manzanas urbanas y de aldeas (89 % de la población): quedan fuera las viviendas rurales dispersas. El censo es de 2024 y la plantación de 2022, y el criterio amplio cuenta la manzana completa aunque solo una parte esté dentro de la distancia, por eso se da el rango con el criterio estricto.",
    "La revisión de trabajos previos no fue exhaustiva.",
])

if REVISION:
    doc.add_heading("12. Revisión externa", 1)
    p(REVISION)

doc.add_page_break()
doc.add_heading("Referencias", 1)
refs = [
    "Gelabert Vadillo, P., Félez Bernal, J., Crespo Pérez, R., Alcasena Urdioz, F. y Rodrigues Mimbrero, M. (2026). Cartografía y análisis de la Interfaz Urbano-Forestal en Chile. Actas del XXI Congreso de la Asociación Española de Teledetección. https://doi.org/10.17398/3101-7177.2.133",
    "Hernández, H. J., González, V., Cortés, L., Hormazábal, C., Muñoz, et al. (2024). Assessing Chile's national expansion of urban and artificial coverages between 2000 and 2022. ISPRS Archives XLVIII-3-2024. https://doi.org/10.5194/isprs-archives-XLVIII-3-2024-643-2024",
    "Hora, B., González-Mathiesen, C., Aravena-Solís, N. y Tapia, T. (2025). Land Cover and Wildfire Risk: A Multi-Buffer Spatial Analysis of the Relationship Between Housing Destruction and Land Cover in Chile's Bío-Bío Region in 2023. Sustainability 17(10), 4416. https://doi.org/10.3390/su17104416",
    "Jaque Castillo, E., Fernández, A., Fuentes Robles, R. y Ojeda, C. G. (2021). Data-based wildfire risk model for Mediterranean ecosystems - case study of the Concepción metropolitan area in central Chile. Natural Hazards and Earth System Sciences 21, 3663-3678. https://doi.org/10.5194/nhess-21-3663-2021",
    "INE (2025). Censo de Población y Vivienda 2024, cartografía censal (manzanas) y diccionario de variables. https://censo2024.ine.gob.cl/resultados/",
    "Ministerio de Educación (2026). Visor comunidades educativas ante amenazas socionaturales. https://emergenciaydesastres.mineduc.cl/visor-comunidades-educativas-ante-amenazas-socionaturales/",
    "McWethy, D. B., Pauchard, A., García, R. A., Holz, A., González, M. E., et al. (2018). Landscape drivers of recent fire activity (2001-2017) in south-central Chile. PLOS ONE 13(8), e0201195. https://doi.org/10.1371/journal.pone.0201195",
    "Sarricolea, P., Serrano-Notivoli, R., Fuentealba, M., Hernández-Mora, M., de la Barrera, F., Smith, P. y Meseguer-Ruiz, Ó. (2020). Recent wildfires in Central Chile: Detecting links between burned areas and population exposure in the wildland urban interface. Science of the Total Environment 706, 135894. https://doi.org/10.1016/j.scitotenv.2019.135894",
    "Schug, F., Bar-Massada, A., Carlson, A. R., Cox, H., et al. (2023). The global wildland-urban interface. Nature 621, 94-99. https://doi.org/10.1038/s41586-023-06320-0",
    "CONAF (2024). Análisis Nacional de Riesgo de Incendios Forestales y Zonas de Interfaz. https://www.conaf.cl/incendios/analisis-nacional-de-riesgo-de-incendios-forestales/ Visor: https://geprif.maps.arcgis.com/apps/dashboards/09532374973a495bb185fcbe55711cd3",
    "Universidad de Chile (2026). MapBiomas Fuego revela 1,2 millones de hectáreas quemadas en 13 años. https://uchile.cl/noticias/244116/mapbiomas-fuego-revela-12-millones-de-hectareas-quemadas-en-13-anos",
    "Corporación Ciudades, con datos de SENAPRED al 6-7-2026, citado en G5 Noticias (10-8-2026). https://g5noticias.cl/2026/08/10/a-cinco-anos-de-la-nueva-institucionalidad-un-40%de-las-comunas-de-chile-siguen-sin-un-plan-formalizado-para-reducir-el-riesgo-de-desastres/",
    "Ley 21.364 (2021), artículo 28. Biblioteca del Congreso Nacional de Chile.",
    "SENAPRED (2026). Visor Chile Preparado. https://senapred.cl/visor-chile-preparado/",
    "NOAA Climate Prediction Center (2026). ENSO Diagnostic Discussion, 8 de octubre de 2026. https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/enso_advisory/ensodisc.shtml",
    "OpenStreetMap (2026). Lugares con nombre (place), consultados el 30-10-2026 vía Overpass API. © colaboradores de OpenStreetMap, licencia ODbL. https://www.openstreetmap.org/copyright",
]
for t in refs:
    par = doc.add_paragraph(t)
    par.paragraph_format.left_indent = Cm(0.6)
    par.paragraph_format.first_line_indent = Cm(-0.6)
    for run in par.runs:
        run.font.size = Pt(8.5)
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.line_spacing = 1.1

doc.add_heading("Cita de los datos de MapBiomas", 2)
for t in (
    "MapBiomas – Colección 1 de la Serie Anual de Mapas de Cobertura y Uso del Suelo de Chile, consultada el 30/10/2026 a través del enlace: https://chile.mapbiomas.org",
    "MapBiomas – Colección 2 de la Serie Anual de Mapas de Cobertura y Uso del Suelo de Chile, consultada el 30/10/2026 a través del enlace: https://chile.mapbiomas.org",
    "MapBiomas – Colección 1 de MapBiomas Fuego Chile (área quemada anual y frecuencia de quema), consultada el 30/10/2026 a través del enlace: https://chile.mapbiomas.org",
):
    par = doc.add_paragraph(t)
    par.paragraph_format.space_after = Pt(3)
    for run in par.runs:
        run.font.size = Pt(8.5)

doc.save("out/Limite_Vivo_MapBiomas_2026.docx")
print("ok")
